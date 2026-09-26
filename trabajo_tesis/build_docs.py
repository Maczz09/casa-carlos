from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


SRC_MATRIX = Path(r"C:\Users\maxmo\Downloads\Título tentativo del proyecto.docx")
SRC_TEMPLATE = Path(r"C:\Users\maxmo\Downloads\Formato Proyecto.docx")
OUT_DIR = Path(r"D:\Hoteles\Casa Carlos\Resultados")
OUT_MATRIX = OUT_DIR / "Título tentativo del proyecto - objetivos revisados.docx"
OUT_CHAPTER = OUT_DIR / "Formato Proyecto - Capítulo I desarrollado.docx"


def set_run_font(run, name="Arial", size=11, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_labeled_paragraph(paragraph, label, body, *, size=11):
    clear_runs(paragraph)
    label_run = paragraph.add_run(label + "\u00a0")
    set_run_font(label_run, size=size, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, size=size, bold=False)


def revise_matrix():
    doc = Document(SRC_MATRIX)
    table = doc.tables[1]
    objectives = [
        "Evaluar si la implementación de un sistema IoT mejora el monitoreo de las condiciones del suelo.",
        "Analizar si la implementación de machine learning mejora el monitoreo de las condiciones del suelo.",
        "Establecer si la implementación de un sistema IoT mejora el monitoreo en el momento de cosecha.",
        "Comprobar si la implementación de machine learning mejora el monitoreo en el momento de cosecha.",
    ]
    hypotheses = [
        "La implementación de un sistema IoT optimizará significativamente el monitoreo de las condiciones del suelo.",
        "La implementación de machine learning potenciará significativamente el monitoreo de las condiciones del suelo.",
        "La implementación de un sistema IoT fortalecerá significativamente el monitoreo en el momento de cosecha.",
        "La implementación de machine learning perfeccionará significativamente el monitoreo en el momento de cosecha.",
    ]
    for i in range(4):
        set_labeled_paragraph(table.cell(i + 1, 1).paragraphs[0], f"OE{i + 1}:", objectives[i], size=11)
        set_labeled_paragraph(table.cell(i + 1, 2).paragraphs[0], f"HE{i + 1}:", hypotheses[i], size=11)
        for c in (1, 2):
            table.cell(i + 1, c).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.core_properties.title = "Sistema IoT para el monitoreo de condiciones del suelo y momento de cosecha"
    doc.core_properties.subject = "Objetivos e hipótesis específicas con verbos revisados"
    doc.save(OUT_MATRIX)


def body_paragraphs(doc):
    return list(doc.paragraphs)


def find_paragraph(doc, exact_text):
    for p in body_paragraphs(doc):
        if p.text.strip() == exact_text:
            return p
    raise ValueError(f"No se encontró el párrafo: {exact_text}")


def delete_between(start, end):
    preserved_sect_pr = None
    preserved_page_break = False
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        for br in node.iter(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                preserved_page_break = True
        p_pr = node.find(qn("w:pPr"))
        if p_pr is not None:
            sect_pr = p_pr.find(qn("w:sectPr"))
            if sect_pr is not None:
                preserved_sect_pr = deepcopy(sect_pr)
        node.getparent().remove(node)
        node = nxt
    return preserved_sect_pr, preserved_page_break


def insert_paragraph_before(anchor, *, label=None, text="", kind="prose", placeholder=False):
    p_xml = OxmlElement("w:p")
    anchor._p.addprevious(p_xml)
    p = Paragraph(p_xml, anchor._parent)
    p.style = "Normal"

    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.keep_together = False
    fmt.keep_with_next = False

    if kind == "prose":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Cm(1.27)
        fmt.line_spacing = 1.5
        fmt.space_after = Pt(6)
    elif kind == "item":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.left_indent = Cm(0.63)
        fmt.first_line_indent = Cm(-0.63)
        fmt.line_spacing = 1.5
        fmt.space_after = Pt(5)
    elif kind == "note":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.63)
        fmt.right_indent = Cm(0.63)
        fmt.line_spacing = 1.15
        fmt.space_after = Pt(6)
    elif kind == "reference":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(1.27)
        fmt.first_line_indent = Cm(-1.27)
        fmt.line_spacing = 1.0
        fmt.space_after = Pt(6)

    if label:
        r = p.add_run(label + "\u00a0")
        set_run_font(r, bold=True)
        r = p.add_run(text)
    else:
        r = p.add_run(text)
    set_run_font(r, italic=(kind == "note"))

    if placeholder:
        r.font.highlight_color = 7  # yellow
        r.bold = True
    return p


def replace_section(doc, start_text, end_text, items):
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    preserved_sect_pr, preserved_page_break = delete_between(start, end)
    inserted = []
    for item in items:
        inserted.append(insert_paragraph_before(end, **item))
    if preserved_sect_pr is not None:
        target = inserted[-1] if inserted else insert_paragraph_before(end, text="")
        target_p_pr = target._p.get_or_add_pPr()
        old_sect_pr = target_p_pr.find(qn("w:sectPr"))
        if old_sect_pr is not None:
            target_p_pr.remove(old_sect_pr)
        target_p_pr.append(preserved_sect_pr)
    if preserved_page_break:
        end.paragraph_format.page_break_before = True


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def develop_chapter():
    doc = Document(SRC_TEMPLATE)

    # Preserve the template's style system; ensure the existing Normal style has
    # a deterministic base for the newly inserted body text.
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)

    description = [
        {
            "kind": "note",
            "label": "Datos de delimitación pendientes:",
            "text": "cultivo y variedad [POR DEFINIR], unidad agrícola y ubicación exacta [POR DEFINIR], y campaña o periodo de evaluación [POR DEFINIR]. Estos datos deben completarse antes de aprobar la versión final del proyecto.",
            "placeholder": True,
        },
        {
            "text": "La producción agrícola depende de la capacidad de conocer, con oportunidad y precisión, el estado del suelo y la evolución del cultivo. A escala mundial, el suelo sostiene aproximadamente el 95 % de los alimentos, pero cerca del 33 % de los suelos ya se encuentra degradado. Esta situación se relaciona con pérdidas de materia orgánica y biodiversidad, desequilibrios de nutrientes y erosión, factores que reducen la productividad y aumentan la vulnerabilidad de los sistemas agrícolas (Organización de las Naciones Unidas para la Alimentación y la Agricultura [FAO], 2025). Por ello, monitorear variables edáficas ya no constituye solo una actividad técnica complementaria, sino una condición para tomar decisiones agrícolas sostenibles y oportunas.",
        },
        {
            "text": "El problema adquiere mayor relevancia por la presión sobre el agua y los insumos. La agricultura concentra más del 70 % de las extracciones mundiales de agua dulce, de modo que una decisión tardía o basada únicamente en observaciones ocasionales puede ocasionar riego ineficiente, fertilización imprecisa y costos innecesarios (FAO, 2020). En la práctica convencional, el productor suele depender de inspecciones manuales, muestreos aislados o análisis de laboratorio que, aunque necesarios como referencia, no siempre ofrecen continuidad temporal. La ausencia de datos frecuentes dificulta identificar cambios en el suelo, comparar tendencias y responder antes de que el estrés del cultivo se vuelva visible.",
        },
        {
            "text": "En el Perú persiste una brecha de tecnificación. La Encuesta Nacional Agropecuaria 2023 informó que solo el 15,8 % de la superficie agrícola empleaba riego tecnificado (Instituto Nacional de Estadística e Informática [INEI], 2024). En Ica, el contexto es especialmente sensible porque la agricultura y la agroexportación dependen de una gestión eficiente del agua subterránea; la Autoridad Nacional del Agua ha señalado la necesidad de recargar y gestionar sosteniblemente los acuíferos para atender a la población y a las actividades agrícolas (Autoridad Nacional del Agua [ANA], 2024). En este escenario, disponer de información continua sobre el suelo puede apoyar decisiones de manejo más responsables, aunque la tecnología no sustituye el análisis agronómico ni la validación de laboratorio.",
        },
        {
            "text": "Los sistemas basados en Internet de las cosas (IoT) ofrecen una alternativa al integrar sensores, un microcontrolador, conectividad, almacenamiento y una interfaz de consulta. Esta arquitectura permite capturar y transmitir lecturas con una frecuencia definida, conservar el historial y emitir alertas cuando los valores se apartan de rangos establecidos. Postolache et al. (2023) explican que los sistemas IoT para evaluación de nutrientes pueden combinar mediciones de humedad, pH, nitrógeno, fósforo y potasio con variables ambientales, pero también enfrentan variabilidad espacial, diferencias entre tipos de suelo, mantenimiento y conectividad. En consecuencia, el valor del sistema no depende solo de “conectar sensores”, sino de asegurar calibración, trazabilidad, continuidad de los datos y una visualización comprensible.",
        },
        {
            "text": "El machine learning puede complementar el monitoreo al reconocer patrones en los datos históricos y producir clasificaciones o estimaciones. Un sistema IoT habilitado con aprendizaje automático puede analizar humedad, temperatura, condiciones ambientales y niveles de nutrientes para generar recomendaciones agrícolas (Islam et al., 2023). Sin embargo, la estimación del momento de cosecha exige definir primero el cultivo, la variedad y un criterio agronómico de madurez que funcione como etiqueta o valor de referencia. La literatura sobre madurez pre-cosecha indica que los modelos suelen requerir fusión de sensores e indicadores específicos, como color, firmeza, contenido de azúcares, acidez o información espectral; por tanto, los datos de N, P y K por sí solos no garantizan una predicción válida del momento de cosecha (Islam et al., 2024).",
        },
        {
            "text": "También existe un riesgo técnico que debe formar parte del problema de investigación: la confiabilidad de los sensores NPK de bajo costo. Baldi et al. (2026) evaluaron sondas comerciales en suelos franco-arcillosos y arenosos y no hallaron correlaciones significativas entre sus lecturas de N, P y K y los análisis químicos de referencia. El hallazgo no invalida todo desarrollo IoT, pero obliga a que el proyecto incluya calibración, comparación con un método de referencia, control de humedad y temperatura, y métricas de error. Sin estas verificaciones, el sistema podría transmitir datos en tiempo real que no representen adecuadamente la condición real del suelo.",
        },
        {
            "text": "En la unidad agrícola [POR DEFINIR] no se cuenta, según el planteamiento disponible, con una solución integrada y validada que registre de forma continua las condiciones del suelo y utilice esos datos para apoyar la estimación del momento de cosecha del cultivo [POR DEFINIR]. Esta carencia limita la disponibilidad de información oportuna, dificulta comparar el comportamiento del suelo durante la campaña y mantiene decisiones sujetas a observaciones discontinuas. Por ello, la investigación propone implementar un sistema IoT con machine learning y evaluar, mediante indicadores verificables, en qué medida mejora el monitoreo de las condiciones del suelo y el apoyo a la determinación del momento de cosecha. La evaluación deberá comparar el proceso actual con el sistema propuesto y reportar exactitud de sensores, disponibilidad y latencia de datos, desempeño del modelo y utilidad de las alertas.",
        },
    ]
    replace_section(doc, "Descripción del Problema", "Formulación del Problema", description)

    replace_section(doc, "Pregunta General", "Preguntas Específicas", [
        {
            "kind": "item",
            "label": "PG:",
            "text": "¿En qué medida la implementación de un sistema IoT y machine learning mejora el monitoreo de las condiciones del suelo y el momento de cosecha?",
        }
    ])

    replace_section(doc, "Preguntas Específicas", "Objetivo de la Investigación", [
        {"kind": "item", "label": "PE1:", "text": "¿En qué medida la implementación de un sistema IoT mejora el monitoreo de las condiciones del suelo?"},
        {"kind": "item", "label": "PE2:", "text": "¿En qué medida la implementación de machine learning mejora el monitoreo de las condiciones del suelo?"},
        {"kind": "item", "label": "PE3:", "text": "¿En qué medida la implementación de un sistema IoT mejora el monitoreo en el momento de cosecha?"},
        {"kind": "item", "label": "PE4:", "text": "¿En qué medida la implementación de machine learning mejora el monitoreo en el momento de cosecha?"},
    ])

    replace_section(doc, "Objetivo General", "Objetivos Específicos", [
        {
            "kind": "item",
            "label": "OG:",
            "text": "Determinar si la implementación de un sistema IoT y machine learning mejora el monitoreo de las condiciones del suelo y el momento de cosecha.",
        }
    ])

    replace_section(doc, "Objetivos Específicos", "Hipótesis", [
        {"kind": "item", "label": "OE1:", "text": "Evaluar si la implementación de un sistema IoT mejora el monitoreo de las condiciones del suelo."},
        {"kind": "item", "label": "OE2:", "text": "Analizar si la implementación de machine learning mejora el monitoreo de las condiciones del suelo."},
        {"kind": "item", "label": "OE3:", "text": "Establecer si la implementación de un sistema IoT mejora el monitoreo en el momento de cosecha."},
        {"kind": "item", "label": "OE4:", "text": "Comprobar si la implementación de machine learning mejora el monitoreo en el momento de cosecha."},
    ])

    replace_section(doc, "Hipótesis General", "Hipótesis Específicas", [
        {
            "kind": "item",
            "label": "HG:",
            "text": "La implementación de un sistema IoT y machine learning mejorará significativamente el monitoreo de las condiciones del suelo y el momento de cosecha.",
        }
    ])

    replace_section(doc, "Hipótesis Específicas", "Justificación e Importancia", [
        {"kind": "item", "label": "HE1:", "text": "La implementación de un sistema IoT optimizará significativamente el monitoreo de las condiciones del suelo."},
        {"kind": "item", "label": "HE2:", "text": "La implementación de machine learning potenciará significativamente el monitoreo de las condiciones del suelo."},
        {"kind": "item", "label": "HE3:", "text": "La implementación de un sistema IoT fortalecerá significativamente el monitoreo en el momento de cosecha."},
        {"kind": "item", "label": "HE4:", "text": "La implementación de machine learning perfeccionará significativamente el monitoreo en el momento de cosecha."},
    ])

    justification = [
        {
            "label": "Justificación práctica.",
            "text": "La investigación busca reducir la dependencia de observaciones aisladas mediante datos continuos y alertas oportunas. Un productor o responsable técnico podrá consultar el historial de las variables del suelo, identificar desviaciones y documentar las condiciones que acompañan el desarrollo del cultivo. Esta trazabilidad puede mejorar la oportunidad de las decisiones y disminuir el tiempo destinado a recopilar manualmente información dispersa.",
        },
        {
            "label": "Justificación tecnológica.",
            "text": "El proyecto integrará sensores, adquisición de datos, comunicación, almacenamiento, procesamiento mediante machine learning y una interfaz de consulta en una arquitectura de software verificable. Su importancia no radica únicamente en construir un prototipo, sino en demostrar interoperabilidad, disponibilidad de datos, tolerancia a fallas de conectividad y capacidad de actualización. El diseño podrá servir como base reutilizable para incorporar sensores adicionales o adaptar el modelo a otro cultivo, siempre después de una nueva calibración y validación.",
        },
        {
            "label": "Justificación metodológica.",
            "text": "La investigación generará un procedimiento de evaluación que confronte las lecturas de los sensores con mediciones de referencia, controle la calidad de los datos y compare el desempeño antes y después de la implementación. Para el componente IoT deberán medirse, como mínimo, error de lectura, porcentaje de datos recibidos, latencia y disponibilidad. Para machine learning se definirán métricas coherentes con la salida: exactitud, precisión, exhaustividad y F1 si se clasifica la madurez; o MAE, RMSE y R² si se estima un valor continuo o los días restantes para la cosecha.",
        },
        {
            "label": "Justificación ambiental y económica.",
            "text": "Un monitoreo mejor fundamentado puede contribuir al uso racional de agua y fertilizantes, especialmente en un territorio con presión hídrica como Ica. La información no reemplazará la recomendación agronómica, pero puede ayudar a evitar intervenciones por rutina o demasiado tardías. En el plano económico, la anticipación de condiciones desfavorables y una estimación mejor documentada de la cosecha pueden reducir pérdidas, reprocesos y costos de supervisión.",
        },
        {
            "label": "Importancia científica y social.",
            "text": "El estudio producirá un conjunto de datos local y evidencia sobre el comportamiento real de sensores y modelos en un cultivo y campaña determinados. Esto es relevante porque los resultados de machine learning dependen del contexto, y un modelo entrenado con datos externos no puede asumirse válido en Ica sin pruebas. La documentación del método, sus errores y sus límites facilitará la replicación académica y permitirá comunicar al usuario final el nivel de confianza de cada recomendación.",
        },
    ]
    replace_section(doc, "Justificación e Importancia", "Limitaciones", [
        {"kind": "prose", "label": item["label"], "text": item["text"]} for item in justification
    ])

    limitations = [
        {
            "text": "La principal limitación actual es la falta de delimitación del cultivo, variedad, unidad agrícola y periodo de estudio. Estas decisiones son indispensables porque determinan los rangos agronómicos, la forma de obtener la verdad de terreno y el significado de “momento de cosecha”. Mientras permanezcan pendientes, el capítulo debe considerarse un borrador técnicamente sustentado y no una versión cerrada.",
        },
        {
            "text": "La exactitud de los sensores, en especial las sondas NPK de bajo costo, puede variar por textura, humedad, temperatura, salinidad, profundidad de instalación, interferencias iónicas y deriva. El estudio no deberá emplear sus lecturas como valores absolutos sin calibrarlas y compararlas con análisis de referencia. Si la validación no alcanza un error aceptable definido previamente, las lecturas solo podrán interpretarse como tendencias relativas.",
        },
        {
            "text": "La conectividad y la alimentación eléctrica pueden provocar pérdida o retraso de datos. La arquitectura deberá contemplar almacenamiento local temporal, reintentos de transmisión, marcas de tiempo, identificación de valores faltantes y monitoreo del estado del nodo. La evaluación se limitará a la tecnología de comunicación y cobertura disponible en la unidad agrícola seleccionada.",
        },
        {
            "text": "El desempeño de machine learning dependerá del número y equilibrio de observaciones etiquetadas. Una sola campaña puede no cubrir toda la variabilidad climática, agronómica y fenológica; por ello, los resultados no deberán generalizarse automáticamente a otras variedades, suelos, fundos o estaciones. La separación de datos de entrenamiento, validación y prueba deberá evitar que mediciones cercanas en el tiempo produzcan una estimación artificialmente optimista.",
        },
        {
            "text": "El momento de cosecha no depende exclusivamente de N, P y K. Según el cultivo, puede requerir humedad, temperatura, grados-día, color, firmeza, sólidos solubles, acidez, imágenes u otros indicadores de madurez. Si el proyecto mantiene únicamente detectores de NPK, deberá presentar el resultado como apoyo al monitoreo y no como determinación agronómica definitiva de la cosecha.",
        },
        {
            "text": "Finalmente, el alcance estará condicionado por presupuesto, disponibilidad de análisis de laboratorio, acceso al campo, eventos meteorológicos, mantenimiento de sensores y participación del especialista agrónomo. Estas restricciones deberán registrarse durante la ejecución para distinguir entre fallas del prototipo, problemas de datos y factores externos al sistema.",
        },
    ]
    replace_section(doc, "Limitaciones", "CAPÍTULO 2", limitations)

    references = [
        "Autoridad Nacional del Agua. (2024, 21 de mayo). Ica: ANA supervisa recarga del acuífero de Villacuri en el sector del parque Golda Meir. https://www.gob.pe/institucion/ana/noticias/958786-ica-ana-supervisa-recarga-del-acuifero-de-villacuri-en-el-sector-del-parque-golda-meir",
        "Baldi, E., Quartieri, M., Chiarelli, G., Tassinari, A., Larocca, G. N., Messini, M., & Toselli, M. (2026). Evaluation of the reliability of in situ, real-time probes for precise determination of nitrogen, phosphorus and potassium in agricultural crops. Horticulturae, 12(7), 898. https://doi.org/10.3390/horticulturae12070898",
        "Instituto Nacional de Estadística e Informática. (2024). Productores agropecuarios: Principales resultados de la Encuesta Nacional Agropecuaria 2023. https://proyectos.inei.gob.pe/iinei/srienaho/Descarga/DocumentosMetodologicos/2023-62/05_PUBLICACION_ENA_2023.pdf",
        "Islam, M., Bijjahalli, S., Fahey, T., Gardi, A., Sabatini, R., & Lamb, D. W. (2024). Destructive and non-destructive measurement approaches and the application of AI models in precision agriculture: A review. Precision Agriculture, 25, 1127-1180. https://doi.org/10.1007/s11119-024-10112-5",
        "Islam, M. R., Oliullah, K., Kabir, M. M., Alom, M., & Mridha, M. F. (2023). Machine learning enabled IoT system for soil nutrients monitoring and crop recommendation. Journal of Agriculture and Food Research, 14, 100880. https://doi.org/10.1016/j.jafr.2023.100880",
        "Organización de las Naciones Unidas para la Alimentación y la Agricultura. (2020). The state of food and agriculture 2020: Overcoming water challenges in agriculture. https://www.fao.org/interactive/state-of-food-agriculture/2020/en/",
        "Organización de las Naciones Unidas para la Alimentación y la Agricultura. (2025, 17 de junio). How healthy soils combat climate change and boost food security. https://www.fao.org/gcf/news/news-detail/how-healthy-soils-combat-climate-change-and-boost-food-security/en",
        "Postolache, S., Sebastião, P., Viegas, V., Postolache, O., & Cercas, F. (2023). IoT-based systems for soil nutrients assessment in horticulture. Sensors, 23(1), 403. https://doi.org/10.3390/s23010403",
    ]
    bibliography = find_paragraph(doc, "BIBLIOGRAFÍA")
    after_bibliography = False
    for paragraph in list(doc.paragraphs):
        if paragraph._p is bibliography._p:
            after_bibliography = True
            continue
        if after_bibliography and not paragraph.text.strip():
            paragraph._p.getparent().remove(paragraph._p)
    # The bibliography is the last content block; append references after it.
    for ref in references:
        p = doc.add_paragraph()
        p.style = "Normal"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(ref)
        set_run_font(r)

    set_update_fields(doc)
    doc.core_properties.title = "Sistema IoT para el monitoreo de condiciones del suelo y momento de cosecha"
    doc.core_properties.subject = "Capítulo I: problema de investigación"
    doc.core_properties.comments = "Borrador académico con delimitaciones pendientes resaltadas."
    doc.save(OUT_CHAPTER)


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    revise_matrix()
    develop_chapter()
    print(OUT_MATRIX)
    print(OUT_CHAPTER)
