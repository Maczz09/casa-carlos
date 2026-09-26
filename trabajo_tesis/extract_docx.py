import json
import sys
import zipfile
from pathlib import Path

from docx import Document


def paragraph_record(p, location):
    text = p.text
    return {
        "location": location,
        "style": p.style.name if p.style else None,
        "text": text,
        "runs": [
            {
                "text": r.text,
                "bold": r.bold,
                "italic": r.italic,
                "underline": r.underline,
                "font": r.font.name,
                "size_pt": r.font.size.pt if r.font.size else None,
            }
            for r in p.runs
        ],
    }


def walk_cell(cell, prefix, out):
    for i, p in enumerate(cell.paragraphs):
        out.append(paragraph_record(p, f"{prefix}/p[{i}]"))
    for ti, table in enumerate(cell.tables):
        walk_table(table, f"{prefix}/table[{ti}]", out)


def walk_table(table, prefix, out):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            walk_cell(cell, f"{prefix}/r[{ri}]/c[{ci}]", out)


def extract(path):
    doc = Document(path)
    paragraphs = [paragraph_record(p, f"body/p[{i}]") for i, p in enumerate(doc.paragraphs)]
    table_paragraphs = []
    for ti, table in enumerate(doc.tables):
        walk_table(table, f"body/table[{ti}]", table_paragraphs)

    sections = []
    hf_paragraphs = []
    for si, section in enumerate(doc.sections):
        sections.append({
            "index": si,
            "width_twips": section.page_width.twips,
            "height_twips": section.page_height.twips,
            "top_margin_twips": section.top_margin.twips if section.top_margin else None,
            "bottom_margin_twips": section.bottom_margin.twips if section.bottom_margin else None,
            "left_margin_twips": section.left_margin.twips if section.left_margin else None,
            "right_margin_twips": section.right_margin.twips if section.right_margin else None,
            "start_type": str(section.start_type),
        })
        for label, part in (("header", section.header), ("footer", section.footer),
                            ("first_header", section.first_page_header), ("first_footer", section.first_page_footer)):
            for pi, p in enumerate(part.paragraphs):
                hf_paragraphs.append(paragraph_record(p, f"section[{si}]/{label}/p[{pi}]"))
            for ti, table in enumerate(part.tables):
                walk_table(table, f"section[{si}]/{label}/table[{ti}]", hf_paragraphs)

    styles = []
    for s in doc.styles:
        if s.type == 1:
            styles.append({
                "name": s.name,
                "font": s.font.name,
                "size_pt": s.font.size.pt if s.font.size else None,
                "bold": s.font.bold,
                "italic": s.font.italic,
            })

    package_parts = []
    with zipfile.ZipFile(path) as zf:
        package_parts = sorted(zf.namelist())

    return {
        "path": str(Path(path).resolve()),
        "paragraphs": paragraphs,
        "table_paragraphs": table_paragraphs,
        "headers_footers": hf_paragraphs,
        "sections": sections,
        "styles": styles,
        "inline_shapes": len(doc.inline_shapes),
        "tables": len(doc.tables),
        "package_parts": package_parts,
    }


if __name__ == "__main__":
    payload = json.dumps(extract(sys.argv[1]), ensure_ascii=False, indent=2)
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(payload, encoding="utf-8")
    else:
        print(payload)
