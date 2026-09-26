from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


def extract(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": paragraph.text,
            }
        )

    tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            rows.append(
                {
                    "row": row_index,
                    "cells": [cell.text for cell in row.cells],
                }
            )
        tables.append({"table": table_index, "rows": rows})

    package = {
        "comments": False,
        "tracked_insertions": 0,
        "tracked_deletions": 0,
        "footnotes": False,
        "endnotes": False,
    }
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        package["comments"] = "word/comments.xml" in names
        package["footnotes"] = "word/footnotes.xml" in names
        package["endnotes"] = "word/endnotes.xml" in names
        xml = archive.read("word/document.xml")
        package["tracked_insertions"] = xml.count(b"<w:ins")
        package["tracked_deletions"] = xml.count(b"<w:del")
        root = ET.fromstring(xml)
        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        tracked = []
        for item in root.findall(f".//{{{word_ns}}}ins"):
            text = "".join(node.text or "" for node in item.findall(f".//{{{word_ns}}}t"))
            tracked.append(
                {
                    "author": item.attrib.get(f"{{{word_ns}}}author"),
                    "date": item.attrib.get(f"{{{word_ns}}}date"),
                    "text": text,
                }
            )
        package["tracked_insertion_details"] = tracked
        if "docProps/app.xml" in names:
            app_root = ET.fromstring(archive.read("docProps/app.xml"))
            package["extended_properties"] = {
                node.tag.rsplit("}", 1)[-1]: node.text for node in app_root
            }

    return {
        "path": str(path),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "package": package,
    }


def main() -> None:
    for raw in sys.argv[1:]:
        path = Path(raw)
        result = extract(path)
        output = Path(__file__).parent / f"{path.stem}.json"
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(output)


if __name__ == "__main__":
    main()
